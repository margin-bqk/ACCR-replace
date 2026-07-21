#!/usr/bin/env python3
"""
ACCR-Replace CLI：批量替换 Excel 中的文本内容。

用法：
    python accr_cli.py \
        --rules 替换规则.xlsx \
        --input 源标题.xlsx \
        --output 输出.xlsx \
        --boundary 6 \
        --col A

输出 Excel：A 列保留原文，B 列写入替换后的文本。
"""

import argparse
import sys
import os
import time

# 确保能 import ACCR_Replace
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from ACCR_Replace.rule_loader import RuleLoader
from ACCR_Replace.ac_engine import ACEngine


def col_letter_to_index(letter):
    """列字母转索引 (A->1, B->2, ..., Z->26, AA->27)"""
    result = 0
    for c in letter.upper():
        result = result * 26 + (ord(c) - ord('A') + 1)
    return result


def col_index_to_letter(index):
    """列索引转字母 (1->A, 2->B, ..., 26->Z, 27->AA)"""
    result = ''
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def main():
    parser = argparse.ArgumentParser(
        description='ACCR-Replace: 基于 AC 自动机的高性能文本替换工具',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python accr_cli.py --rules rules.xlsx --input titles.xlsx --output result.xlsx
  python accr_cli.py --rules rules.xlsx --input titles.xlsx --output result.xlsx --boundary 6 --col A --sheet Sheet1
        """
    )

    parser.add_argument('--rules', required=True, help='规则 Excel 文件路径')
    parser.add_argument('--input', required=True, help='输入 Excel 文件路径（待替换文本）')
    parser.add_argument('--output', required=True, help='输出 Excel 文件路径')
    parser.add_argument('--boundary', type=int, default=6,
                        help='词边界阈值：≤N 字符的规则启用边界检查（默认 6）')
    parser.add_argument('--col', default='A',
                        help='待替换文本所在列（默认 A）')
    parser.add_argument('--sheet', default=None,
                        help='输入文件的工作表名（默认第一个）')
    parser.add_argument('--rules-sheet', default='Sheet1',
                        help='规则文件的工作表名（默认 Sheet1）')
    parser.add_argument('--dry-run', action='store_true',
                        help='只预览不写文件，显示前 20 条替换结果')
    parser.add_argument('--stats', action='store_true',
                        help='显示详细统计信息')

    args = parser.parse_args()

    # 检查文件存在
    if not os.path.exists(args.rules):
        print(f"错误: 规则文件不存在: {args.rules}")
        sys.exit(1)
    if not os.path.exists(args.input):
        print(f"错误: 输入文件不存在: {args.input}")
        sys.exit(1)

    # Step 1: 加载并清洗规则
    print("=" * 60)
    print("Step 1: 加载规则")
    print("=" * 60)

    t0 = time.time()
    loader = RuleLoader(args.rules, sheet_name=args.rules_sheet)
    rules = loader.load()
    strict_rules = loader.strict_rules
    t_rules = time.time() - t0

    print(f"规则加载完成: {len(rules)} 条规则 ({t_rules:.3f}s)")
    if strict_rules:
        print(f"完全边界规则: {len(strict_rules)} 条")

    if args.stats:
        loader.report()

    # Step 2: 构建 AC 自动机
    print("\n" + "=" * 60)
    print("Step 2: 构建 AC 自动机")
    print("=" * 60)

    t0 = time.time()
    engine = ACEngine(rules, boundary_threshold=args.boundary, strict_rules=strict_rules)
    t_build = time.time() - t0

    print(f"AC 自动机构建完成 ({t_build:.3f}s)")
    print(f"词边界阈值: {args.boundary} 字符")

    # Step 3: 读取输入 Excel
    print("\n" + "=" * 60)
    print("Step 3: 读取输入文件")
    print("=" * 60)

    import openpyxl

    wb = openpyxl.load_workbook(args.input)
    sheet_name = args.sheet if args.sheet else wb.sheetnames[0]
    ws = wb[sheet_name]

    src_col = col_letter_to_index(args.col)
    dst_col = src_col + 1
    dst_letter = col_index_to_letter(dst_col)

    # 收集所有待处理文本
    rows_data = []
    for row_idx in range(2, ws.max_row + 1):  # 跳过表头
        cell = ws.cell(row=row_idx, column=src_col)
        value = cell.value
        if value is not None:
            text = str(value).strip()
            if text:
                rows_data.append({
                    'row': row_idx,
                    'original': text,
                })

    total_rows = len(rows_data)
    print(f"工作表: {sheet_name}")
    print(f"源列: {args.col} (第 {src_col} 列)")
    print(f"目标列: {dst_letter} (第 {dst_col} 列)")
    print(f"待处理行数: {total_rows}")

    if total_rows == 0:
        print("没有待处理的数据")
        sys.exit(0)

    # Step 4: 批量替换
    print("\n" + "=" * 60)
    print("Step 4: 执行替换")
    print("=" * 60)

    t0 = time.time()
    changed_count = 0
    match_count = 0
    preview_items = []

    for i, item in enumerate(rows_data):
        result = engine.match_and_replace(item['original'])
        item['replaced'] = result['replaced']
        item['changed'] = result['changed']
        item['matches'] = result['matches']

        if result['changed']:
            changed_count += 1
            match_count += len(result['matches'])

        # 收集预览数据
        if args.dry_run and len(preview_items) < 20:
            preview_items.append(item)

        # 进度显示
        if (i + 1) % 10000 == 0 or (i + 1) == total_rows:
            elapsed = time.time() - t0
            speed = (i + 1) / elapsed if elapsed > 0 else 0
            print(f"  进度: {i + 1}/{total_rows} ({speed:.0f} 行/秒)")

    t_replace = time.time() - t0

    print(f"\n替换完成 ({t_replace:.3f}s)")
    print(f"总耗时: {t_rules + t_build + t_replace:.3f}s")
    print(f"处理速度: {total_rows / t_replace:.0f} 行/秒")
    print(f"有替换的行: {changed_count}/{total_rows} ({changed_count / total_rows * 100:.1f}%)")
    print(f"总匹配次数: {match_count}")

    # Dry run: 只显示预览
    if args.dry_run:
        print("\n" + "=" * 60)
        print("Dry Run 预览（前 20 条）")
        print("=" * 60)
        for item in preview_items:
            status = "→" if item['changed'] else "="
            print(f"\n  行{item['row']}:")
            print(f"    原文: {item['original']}")
            print(f"    {status} 结果: {item['replaced']}")
            if item['matches']:
                for start, end, pattern, replacement in item['matches']:
                    matched_text = item['original'][start:end]
                    action = "删除" if replacement == '' else f"→ \"{replacement}\""
                    print(f"      [{start}:{end}] \"{matched_text}\" (规则: \"{pattern}\") {action}")
        print(f"\n共 {total_rows} 行，使用 --output 写入完整结果")
        return

    # Step 5: 写入输出 Excel
    print("\n" + "=" * 60)
    print("Step 5: 写入输出文件")
    print("=" * 60)

    # 写入表头
    ws.cell(row=1, column=dst_col, value=f"{args.col}_replaced")

    for item in rows_data:
        ws.cell(row=item['row'], column=dst_col, value=item['replaced'])

    wb.save(args.output)
    print(f"已保存: {args.output}")

    # 统计摘要
    if args.stats:
        print("\n" + "=" * 60)
        print("统计摘要")
        print("=" * 60)

        # 按规则统计命中次数
        rule_hits = {}
        for item in rows_data:
            for start, end, pattern, replacement in item['matches']:
                rule_hits[pattern] = rule_hits.get(pattern, 0) + 1

        print("\n命中 Top 20 的规则:")
        for pattern, count in sorted(rule_hits.items(), key=lambda x: -x[1])[:20]:
            replacement = rules.get(pattern, '')
            action = "删除" if replacement == '' else f"→ \"{replacement}\""
            print(f"  {count:5d}x  \"{pattern}\" {action}")

        never_hit = set(rules.keys()) - set(rule_hits.keys())
        if never_hit:
            print(f"\n未命中的规则: {len(never_hit)} 条")

    print("\n完成!")


if __name__ == "__main__":
    main()
